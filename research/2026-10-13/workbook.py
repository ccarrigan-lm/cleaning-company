import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
from pathlib import Path

D = Path(__file__).parent
M = json.load(open(D / 'matrix.json'))
sl = json.load(open(D / 'scrub_lou_model.json'))
ma = json.load(open(D / 'maids_in_apron_model.json'))
sc = json.load(open(D / 'super_cleaning_model.json'))
RRM = json.load(open(D / 'rr_maids_model.json'))

F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF')
HF = PatternFill('solid', fgColor='1F4E78')
B = Font(name=F, bold=True)
N = Font(name=F)
BLUE = Font(name=F, color='0000FF')
YEL = PatternFill('solid', fgColor='FFF2CC')
GRY = PatternFill('solid', fgColor='F2F2F2')
TITLE = Font(name=F, bold=True, size=14)
thin = Side(style='thin', color='BFBFBF')
USD = '$#,##0.00;($#,##0.00);-'
USD0 = '$#,##0;($#,##0);-'
PCT = '0.0%'

wb = Workbook()

def header(ws, row, cols, start=1):
    for i, c in enumerate(cols):
        cell = ws.cell(row=row, column=start + i, value=c)
        cell.font, cell.fill = H, HF
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def fontall(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None,
                              size=c.font.size if c.font else 11, italic=c.font.italic if c.font else False)

# ------------------------------------------------------------------ Summary
ws = wb.active
ws.title = 'Summary'
ws['A1'] = 'Louisville Residential Cleaning – Competitor Price Matrix'
ws['A1'].font = TITLE
ws['A2'] = 'Captured 13 Oct 2026 · ZIP 40217 · prices pre-tax unless noted (KY sales tax 6% is added by all three online-booking competitors)'
lines = [
    ('What is in this workbook', None),
    ('Comparison – Standard', 'Same 19 home profiles priced by every competitor (one-time standard clean), with market low / median / high and your price at a chosen position.'),
    ('Comparison – Recurring', 'Weekly / every-2-weeks / every-4-weeks: first-visit and ongoing per-visit price by competitor.'),
    ('Comparison – Deep & Move', 'One-time deep clean and move-in/out for the same profiles.'),
    ('Competitor Rules', 'How each company actually builds its price (base, per-bedroom, per-bath, sq ft bands, discounts, add-ons).'),
    ('Extras', 'Add-on prices side by side.'),
    ('Published Prices', 'Price lists / ranges that other Louisville companies publish on their websites (no instant calculator).'),
    ('Full Matrix', f"Every combination each competitor's online form can produce: {len(M['full']):,} rows. Filter by company / service / frequency / size."),
    (None, None),
    ('How the data was collected', None),
    ('Scrub Lou', 'Drove the live BookingKoala quote form through all 5,220 combinations (3 services × 4 frequencies × 5 bedrooms × 9 bathrooms × 9–11 sq ft bands) and read the Booking Summary. Every captured price is reproduced exactly by the rules on Competitor Rules.'),
    ('Maids in Apron', 'Same method on their BookingKoala form: 7,056 combinations (6 frequencies × 7 bedrooms × 14 bathrooms × 12 sq ft bands), plus the deep/move add-ons and extras. Rules reproduce every first-visit total exactly.'),
    ('Super Cleaning Service', "Read the price table from their Launch27 booking form's own data feed, then confirmed on the live form (e.g. 3 bd / 2 ba / 1,500–1,999 sq ft = $255.00 + tax)."),
    ('Safety', 'No contact details were entered and nothing was booked or submitted on any site.'),
    (None, None),
    ('Key findings', None),
    ('1', 'Standard one-time clean for a typical 3 bd / 2 ba / ~1,700 sq ft home: $219–$255 at the three instant-quote companies, $220 at Stitch, ~$230 at Cleanoria; Rocking Robins charges $420 for a first-time basic clean. Market median ≈ $240 before tax.'),
    ('2', 'Recurring discounts differ a lot: Super Cleaning 10/25/35% off (4-week/2-week/weekly) including the first visit; Scrub Lou 0/15/20%; Maids in Apron charges full price for the first visit and only $10/$15/$20 off later visits.'),
    ('3', 'Deep cleans run ~1.7× a standard clean and move-in/out ~1.9× (median across competitors), but the spread is wide (1.3×–2.6×).'),
    ('4', 'Scrub Lou prices mostly on square footage (+$20 per 500 sq ft band) with small bed/bath adders (+$10 each). Super Cleaning charges heavily per bathroom ($35 full / $20 half). Maids in Apron jumps steeply for 4+ bathrooms and 3+ bedrooms.'),
    ('5', 'Oven / fridge add-ons: $45 each at Scrub Lou and Maids in Apron vs $65 at Super Cleaning.'),
    ('Round 1', 'See Round 1 Tracker: 4 of 10 companies give instant online prices (Scrub Lou, Maids in Apron, Super Cleaning, Rocking Robins); LouBelle publishes $40/hr per tech (2-tech minimum); 5 need a quote request with contact details (awaiting your OK).'),
    ('Hourly benchmarks', 'Independent cleaners: Care.com average posted $19/hr (listings $15–$40/hr); Taskrabbit "from $40"; Homeaglow ≈$25/hr + $59/mo membership; Thumbtack pros ~$60/hr; LouBelle $80/hr per 2-person crew; Rocking Robins prices on $60 per labor hour.'),
    ('6', f"A simple market formula fits the median one-time standard price within ±${M['fit']['mean_abs_error']:.0f} on average: ${M['fit']['intercept']:.0f} + ${M['fit']['per_1000_sqft']:.0f} per 1,000 sq ft + ${M['fit']['per_bedroom']:.0f}/bedroom + ${M['fit']['per_bathroom']:.0f}/bathroom (R² {M['fit']['r2']})."),
]
r = 4
for a, b in lines:
    if a and b is None:
        ws.cell(row=r, column=1, value=a).font = Font(name=F, bold=True, size=12)
    elif a:
        ws.cell(row=r, column=1, value=a).font = B
        c = ws.cell(row=r, column=2, value=b)
        c.font, c.alignment = N, Alignment(wrap_text=True, vertical='top')
    r += 1
widths(ws, [26, 120])

# ------------------------------------------------------------------ Comparison – Standard
comp = M['comparison']
names = ['Scrub Lou', 'Maids in Apron', 'Super Cleaning Service']
names5 = names + ['Rocking Robins (R&R) Maids', 'Cleanoria (Thumbtack midpoint)']
from openpyxl.utils import get_column_letter as L
ws = wb.create_sheet('Comparison – Standard')
ws['A1'] = 'Standard clean, one time – price before tax'
ws['A1'].font = TITLE
ws['A2'] = 'Your position vs market median:'
ws['A2'].font = B
ws['D2'] = 0.0
ws['D2'].font, ws['D2'].fill, ws['D2'].number_format = BLUE, YEL, '+0%;-0%;0%'
ws['E2'] = '← edit (e.g. -5% to undercut, +10% for premium). "Your price" rounds to the nearest $5.'
ws['E2'].font = Font(name=F, italic=True)
cols = ['Home profile', 'Beds', 'Baths', 'Sq ft'] + names5 + ['Stitch Cleaning (published)', 'Market low', 'Market median', 'Market high', 'Spread (high−low)', 'Your price']
header(ws, 4, cols)
std = [c for c in comp if c['service'] == 'standard' and c['frequency'] == 'one_time']
c0, c1 = 5, 4 + len(names5) + 1          # competitor columns incl. Stitch
cLow, cMed, cHigh, cSpr, cYou = c1 + 1, c1 + 2, c1 + 3, c1 + 4, c1 + 5
STD_MED_COL = L(cMed)
for i, c in enumerate(std):
    rr = 5 + i
    vals = [c['profile'], c['bedrooms'], c['bathrooms'], c['sqft']] + [c[f'{n} first visit'] for n in names5] + [c['Stitch Cleaning (published, one-time)']]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=rr, column=j, value=v)
        cell.font = N
        if j >= 5:
            cell.number_format = USD
    rng_ = f'{L(c0)}{rr}:{L(c1)}{rr}'
    ws.cell(row=rr, column=cLow, value=f'=MIN({rng_})')
    ws.cell(row=rr, column=cMed, value=f'=MEDIAN({rng_})')
    ws.cell(row=rr, column=cHigh, value=f'=MAX({rng_})')
    ws.cell(row=rr, column=cSpr, value=f'={L(cHigh)}{rr}-{L(cLow)}{rr}')
    ws.cell(row=rr, column=cYou, value=f'=MROUND({L(cMed)}{rr}*(1+$D$2),5)')
    for j in range(cLow, cYou + 1):
        ws.cell(row=rr, column=j).number_format = USD
        ws.cell(row=rr, column=j).font = N
    ws.cell(row=rr, column=cYou).font = B
    ws.cell(row=rr, column=cMed).fill = GRY
last = 4 + len(std)
ws.cell(row=last + 2, column=1, value='Notes').font = B
notes = [
    'Each competitor is priced at the sq ft band that contains the profile sq ft. Bathrooms like 2.5 = 2 full + 1 half for Super Cleaning.',
    'Stitch Cleaning publishes flat bed/bath prices only (no sq ft; half baths not priced) – https://stitchcleaning.com/louisville-cleaning-services-pricing/',
    'Market low/median/high include Stitch where it publishes a price.',
    'Rocking Robins (R&R): online price is for the FIRST clean ("1st Time Basic"), average condition, 1 story, no basement; only offered for their listed bed/bath layouts.',
    'Cleanoria: midpoint of the price range the pro publishes on Thumbtack; only for listed bed/bath + sq ft combinations.',
]
for k, t in enumerate(notes):
    ws.cell(row=last + 3 + k, column=1, value=t).font = Font(name=F, italic=True)
widths(ws, [28, 7, 7, 9, 13, 15, 16, 16, 16, 16, 12, 14, 12, 14, 12])
ws.freeze_panes = 'E5'

# ------------------------------------------------------------------ Comparison – Recurring
ws = wb.create_sheet('Comparison – Recurring')
ws['A1'] = 'Standard clean on a recurring schedule – per-visit price before tax'
ws['A1'].font = TITLE
ws['A2'] = 'First visit = what the customer pays for the first clean; Ongoing = each later visit.'
ws['A2'].font = Font(name=F, italic=True)
cols = ['Frequency', 'Home profile', 'Beds', 'Baths', 'Sq ft']
for n in names:
    cols += [f'{n} – first', f'{n} – ongoing']
cols += ['Ongoing median', 'Ongoing vs one-time median']
header(ws, 4, cols)
FREQ_L = {'weekly': 'Weekly', 'biweekly': 'Every 2 weeks', 'every_4_weeks': 'Every 4 weeks'}
rr = 5
onemed = {c['profile']: c['market median'] for c in std}
for fk in ['weekly', 'biweekly', 'every_4_weeks']:
    for c in [x for x in comp if x['service'] == 'standard' and x['frequency'] == fk]:
        vals = [FREQ_L[fk], c['profile'], c['bedrooms'], c['bathrooms'], c['sqft']]
        for n in names:
            vals += [c[f'{n} first visit'], c[f'{n} recurring visit']]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=rr, column=j, value=v)
            cell.font = N
            if j >= 6:
                cell.number_format = USD
        ws.cell(row=rr, column=12, value=f'=MEDIAN(G{rr},I{rr},K{rr})').number_format = USD
        ws.cell(row=rr, column=13, value=f"=L{rr}/'Comparison – Standard'!{STD_MED_COL}{5 + [s['profile'] for s in std].index(c['profile'])}-1").number_format = '+0%;-0%;0%'
        ws.cell(row=rr, column=12).font = N
        ws.cell(row=rr, column=13).font = N
        rr += 1
widths(ws, [14, 28, 7, 7, 9, 12, 12, 12, 12, 12, 12, 13, 14])
ws.freeze_panes = 'F5'

# ------------------------------------------------------------------ Comparison – Deep & Move
ws = wb.create_sheet('Comparison – Deep & Move')
ws['A1'] = 'One-time deep clean and move-in/out – price before tax'
ws['A1'].font = TITLE
cols = ['Service', 'Home profile', 'Beds', 'Baths', 'Sq ft'] + names5 + ['Market low', 'Market median', 'Market high', 'Median ÷ standard median']
header(ws, 3, cols)
d0, d1 = 6, 5 + len(names5)
rr = 4
for svc, lab in [('deep', 'Deep clean'), ('move_in_out', 'Move-in / move-out')]:
    for c in [x for x in comp if x['service'] == svc]:
        vals = [lab, c['profile'], c['bedrooms'], c['bathrooms'], c['sqft']] + [c[f'{n} first visit'] for n in names5]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=rr, column=j, value=v)
            cell.font = N
            if j >= 6:
                cell.number_format = USD
        rng_ = f'{L(d0)}{rr}:{L(d1)}{rr}'
        ws.cell(row=rr, column=d1 + 1, value=f'=MIN({rng_})').number_format = USD
        ws.cell(row=rr, column=d1 + 2, value=f'=MEDIAN({rng_})').number_format = USD
        ws.cell(row=rr, column=d1 + 3, value=f'=MAX({rng_})').number_format = USD
        si = 5 + [s_['profile'] for s_ in std].index(c['profile'])
        ws.cell(row=rr, column=d1 + 4, value=f"={L(d1 + 2)}{rr}/'Comparison – Standard'!{STD_MED_COL}{si}").number_format = '0.00"×"'
        for j in range(d1 + 1, d1 + 5):
            ws.cell(row=rr, column=j).font = N
        rr += 1
ws.cell(row=rr + 1, column=1, value='Maids in Apron and Super Cleaning sell deep / move-in-out as add-ons on top of the standard price; the totals here include the add-on.').font = Font(name=F, italic=True)
widths(ws, [18, 28, 7, 7, 9, 13, 15, 16, 16, 16, 12, 14, 12, 14])
ws.freeze_panes = 'F4'

# ------------------------------------------------------------------ Competitor Rules
ws = wb.create_sheet('Competitor Rules')
ws['A1'] = 'How each competitor builds its price (pre-tax)'
ws['A1'].font = TITLE
r = 3
def block(title, rows, hdr):
    global r
    ws.cell(row=r, column=1, value=title).font = Font(name=F, bold=True, size=12)
    r += 1
    header(ws, r, hdr)
    r += 1
    for row in rows:
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = N
            if isinstance(v, (int, float)) and j > 1:
                c.number_format = USD0
        r += 1
    r += 1

ws.cell(row=r, column=1, value='SCRUB LOU — price = sq ft band base + bedroom adder + bathroom adder, × frequency multiplier').font = B; r += 1
block('Sq ft band base by service', [[s, *v['sqft_base'], *[None] * (11 - len(v['sqft_base']))] for s, v in sl['services'].items()],
      ['Service'] + sl['services']['Routine Cleaning']['sqft'])
block('Bedroom / bathroom adders', [['Bedrooms ' + str(b), a] for b, a in zip(sl['bedrooms'], sl['bedroom_add'])] +
      [['Bathrooms ' + str(b), a] for b, a in zip(sl['bathrooms'], sl['bathroom_add'])], ['Option', 'Adds'])
block('Frequency', [[k, f'{(1 - v):.0%} off every visit'] for k, v in sl['frequency_multiplier'].items()], ['Frequency', 'Discount'])

ws.cell(row=r, column=1, value='MAIDS IN APRON — first visit = sq ft band + bedroom adder + bathroom adder (+ deep/move add-on); ongoing visits = first − flat discount').font = B; r += 1
block('Sq ft band base', [[s, b] for s, b in zip(ma['sqft'], ma['sqft_base'])], ['Sq ft band', 'Base'])
block('Bedroom adders and deep / move add-ons', [[b, a, ma['deep_clean_addon_by_bedrooms'][str(b)], ma['move_in_out_addon_by_bedrooms'][str(b)]] for b, a in zip(ma['bedrooms'], ma['bedroom_add'])],
      ['Bedrooms', 'Adds', 'Deep clean add-on', 'Move in/out add-on'])
block('Bathroom adders', [[b, a] for b, a in zip(ma['bathrooms'], ma['bathroom_add'])], ['Bathrooms', 'Adds'])
block('Frequency', [[k, v] for k, v in ma['recurring_visit_discount'].items()], ['Frequency', 'Ongoing-visit discount (first visit is full price)'])

ws.cell(row=r, column=1, value='SUPER CLEANING SERVICE — price = bedroom base + sq ft band + $35/full bath + $20/half bath; % discount applies to every visit including the first').font = B; r += 1
block('Bedroom base and add-ons', [[b, v, sc['deep_clean_addon_first_visit'][b], sc['move_in_out_addon'][b]] for b, v in sc['bedroom_base'].items()],
      ['Bedrooms', 'Base', 'Deep clean add-on', 'Move in/out add-on'])
block('Sq ft band adders', [[s, a] for s, a in zip(sc['sqft'], sc['sqft_add'])], ['Sq ft band', 'Adds'])
block('Frequency', [[k, f'{v}% off'] for k, v in sc['frequency_discount_pct'].items()], ['Frequency', 'Discount'])
ws.cell(row=r, column=1, value='Hourly option: $65/hr per cleaner, 1–2 cleaners, 4–8 hours.').font = N
widths(ws, [26] + [14] * 11)

# ------------------------------------------------------------------ Extras
ws = wb.create_sheet('Extras')
ws['A1'] = 'Add-on prices (pre-tax, one-time)'
ws['A1'].font = TITLE
header(ws, 3, ['Add-on', 'Scrub Lou', 'Maids in Apron', 'Super Cleaning Service', 'Notes'])
ex = [
    ('Inside oven', 45, 45, 65, ''),
    ('Inside fridge', 45, 45, 65, 'Super Cleaning: "Fridge and Freezer"'),
    ('Interior windows (per window)', 4, 5, 8, 'Scrub Lou: glass only $4, glass & tracks $8'),
    ('Window blinds (per unit)', None, 5, None, ''),
    ('Inside cabinets', None, 5, None, 'Maids in Apron: per unit; only offered for some selections'),
    ('Basement', 'custom quote', 120, None, 'Maids in Apron: finished basement'),
    ('Extra kitchen', None, 90, None, ''),
    ('Pets', None, 10, None, ''),
    ('Laundry & folding', None, 65, None, ''),
    ('Dishes', None, 65, None, ''),
    ('House organization', None, 375, None, ''),
    ('Green cleaning', None, 0, None, ''),
    ('Vacuum sectional / large couch', None, None, 15, ''),
    ('Vacuum small couch', None, None, 10, ''),
    ('Additional time (per 30 min)', None, None, 32.5, 'Not discounted for recurring'),
    ('R&R: windows $5 ea, blinds $20 ea, fridge/freezer $60', None, None, None, 'Rocking Robins (R&R) Maids add-ons'),
    ('Exclude a room (partial clean)', -10, None, None, 'Scrub Lou: −$10 per bedroom / full bath / kitchen / living room, −$5 per half bath'),
]
for i, row in enumerate(ex):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=4 + i, column=j, value=v)
        c.font = N
        if isinstance(v, (int, float)):
            c.number_format = USD
ws.cell(row=5 + len(ex), column=1, value='Recurring discounts also apply to add-ons at Scrub Lou and Super Cleaning.').font = Font(name=F, italic=True)
widths(ws, [32, 14, 16, 22, 70])

# ------------------------------------------------------------------ Published Prices
ws = wb.create_sheet('Published Prices')
ws['A1'] = 'Prices published on other Louisville companies’ websites (no instant calculator captured)'
ws['A1'].font = TITLE
header(ws, 3, ['Company', 'Home size / item', 'Standard', 'Deep', 'Move-in/out', 'Other', 'Source'])
pub = [
    ('Stitch Cleaning', 'Flat price by bed/bath (see Comparison – Standard)', '1bd/1ba $150 … 5bd/5ba $400', '', '', '10% off first booking; up to 15% off weekly/biweekly', 'https://stitchcleaning.com/louisville-cleaning-services-pricing/'),
    ('Your Cleaning Krew', '1BR / apartment', '$86–$150+', '$190–$240+', '$230–$280+', '', 'https://yourcleaningkrew.com/louisville/house-cleaning-prices/'),
    ('Your Cleaning Krew', '2–3BR', '$120–$200+', '$220–$320+', '$260–$400+', '', 'https://yourcleaningkrew.com/louisville/house-cleaning-prices/'),
    ('Your Cleaning Krew', '4+BR', '$180–$300+', '$300–$450+', '$350–$550+', '$60/hr one cleaner (2-hr min); $90/hr two cleaners', 'https://yourcleaningkrew.com/louisville/house-cleaning-prices/'),
    ('R&R Maids', 'Smaller homes / 2,000–3,000 sq ft / larger', '$120–$180 / $180–$250 / $250+ (recurring)', '$400–$700 (most $480–$600)', '$450–$800', '', 'https://rrmaids.com/louisville-house-cleaning-prices/'),
    ('Busy Brooms', 'Studio/1BR', '$100–$140', '$170–$240', '$190–$260', 'May be market ranges, not own rates', 'https://busybroomsky.com/house-cleaning-cost-louisville-ky/'),
    ('Busy Brooms', '2BR / 1,200 sq ft', '$130–$180', '$210–$290', '$240–$330', '', 'https://busybroomsky.com/house-cleaning-cost-louisville-ky/'),
    ('Busy Brooms', '3BR / 1,800 sq ft', '$160–$220', '$260–$350', '$300–$400', '', 'https://busybroomsky.com/house-cleaning-cost-louisville-ky/'),
    ('Busy Brooms', '4BR / 2,500 sq ft', '$220–$320', '$340–$450', '$380–$500', 'Recurring 10–20% off; pets $10–$30', 'https://busybroomsky.com/house-cleaning-cost-louisville-ky/'),
    ('Hello Cleaners', '“From” prices', 'from $160', 'from $200', 'move-out from $190; move-in from $180', 'National operator', 'https://hellocleaners.us/locations/kentucky/louisville/'),
    ('Super Cleaning Service', 'Up to 4bd / 3ba / 2,999 sq ft (own site)', '$160–$400', '$200–$665', '$395–$695', 'Matches the calculator rules', 'https://supercleaningservicelouisville.com/services/'),
]
for i, row in enumerate(pub):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=4 + i, column=j, value=v)
        c.font = N
        c.alignment = Alignment(wrap_text=True, vertical='top')
r = 6 + len(pub)
ws.cell(row=r, column=1, value='Lead-form only (no price online): The Maids (Southern Louisville), MaidPro Greater Louisville, 502 Star Services, Deep Blue Clean, Red Feather Cleaning. Next to automate: R&R Maids, Your Cleaning Krew and Louisville Cleaning Pros online calculators.').font = Font(name=F, italic=True)
widths(ws, [24, 34, 26, 22, 22, 34, 60])

# ------------------------------------------------------------------ Round 1 Tracker
ws = wb.create_sheet('Round 1 Tracker', 1)
ws['A1'] = 'Round 1 quote-request tracker – 10 companies + platform sample (checked 13 Oct 2026, ZIP 40217)'
ws['A1'].font = TITLE
ws['A2'] = 'Status key: CAPTURED = full online pricing in this workbook · PUBLISHED = price list/rate on their site · NEEDS REQUEST = price only after submitting contact details (awaiting your OK) · NO PRICE = contact them directly'
ws['A2'].font = Font(name=F, italic=True)
hdr = ['#', 'Company / provider', 'Type', 'Quote route (URL)', 'How price is given', 'Status', 'What we captured', 'Contact details required to see price?', 'Serves 40217?', 'Next action', 'Promo vs normal price', 'Phone']
header(ws, 4, hdr)
T = [
 (1, 'Scrub Lou', 'Local', 'https://scrublou.bookingkoala.com/booknow', 'Instant online calculator (BookingKoala)', 'CAPTURED', 'All 5,220 combinations; extras; exclusions', 'No', 'Yes', 'Compare with your previous Scrub Lou bill', 'Thumbtack listing says "starting at $150"; recurring 15–20% off', ''),
 (2, 'Maids in Apron', 'Local', 'https://maidsinapron.bookingkoala.com/booknow', 'Instant online calculator (BookingKoala)', 'CAPTURED', 'All 7,056 combinations; deep/move add-ons; 12 extras', 'No', 'Yes', '—', 'Recurring: $10–$20 off ongoing visits only', '(502) 202-7542'),
 (3, 'Super Cleaning Service Louisville', 'Local', 'https://supercleaning.launch27.com/?w_cleaning', 'Instant online calculator (Launch27)', 'CAPTURED', 'Full price table + checklist page; spot-checked', 'No', 'Yes', 'Compare scope using their cleaning checklist page', 'Craigslist ads advertise "up to 45% off"; site discounts 10/25/35% recurring', '(502) 230-6735'),
 (4, 'Maid in the Ville', 'Local', 'https://www.maidintheville.com/request-an-estimate', 'Estimate request form only', 'NEEDS REQUEST', 'No prices published', 'Yes – name, email, phone, address', 'Serves Louisville incl. Prospect, St. Matthews, Germantown', 'Submit estimate request with research email (needs your OK) or call', '—', '502-553-9575'),
 (5, 'LouBelle Cleaning', 'Local', 'https://www.loubellecleaning.com/contact-quote', 'Published hourly rate + phone/in-home estimate', 'PUBLISHED', '$40/hr per technician, minimum 2 technicians (= $80/hr per crew-hour)', 'Yes for exact quote', 'Louisville / Kentuckiana', 'Call 502-641-2507 for estimate (manual)', '—', '502-641-2507'),
 (6, 'Rocking Robins Maids (R&R Maids)', 'Local', 'https://rrmaids.com/book-house-cleaning/', 'Instant online calculator (hours × $60) + $50 deposit to book', 'CAPTURED', 'Full formula: layout, sq ft, condition, stories, basement, basic vs deep, add-ons', 'No (price shows before contact fields)', 'Louisville + Southern Indiana', 'Ask how recurring (after first clean) is priced', 'Online price is first clean only; frequency does not change it', '(502) 488-8685'),
 (7, 'The Cleaning Authority – Louisville', 'National', 'https://estimate.thecleaningauthority.com/', 'Online estimate after contact step', 'NEEDS REQUEST', 'Confirmed ZIP 40217 is served; step 1 asks name, phone, email, address before pricing', 'Yes – name, phone, email, street address', 'Yes (confirmed)', 'Submit step 1 with research email (needs your OK) – price shows on step 2', '—', '502.337.3518'),
 (8, 'MaidPro – Greater Louisville', 'National', 'https://www.maidpro.com/greater-louisville/estimate', 'Estimate request form (Salesforce), call back', 'NEEDS REQUEST', '49-point checklist; promos captured', 'Yes – name, address, email, mobile', 'Yes', 'Submit request (needs OK); record normal price and promo separately', 'Promo: $25 off 1st, 3rd and 5th cleans (new customers); "Save $75" code 75OFF', '(502) 254-9979'),
 (9, 'Merry Maids – Louisville', 'National', 'https://merrymaids.com/home-cleaning-quote', 'Multi-step quote starting with street address', 'NEEDS REQUEST', 'No online prices or offers listed', 'Yes – starts with street address', 'Yes (1807 Cargo Ct, 40299)', 'Submit quote flow (needs OK) or call', 'Offers page lists none', '502.369.0887'),
 (10, 'The Maids – Louisville', 'National', 'https://www.maids.com/estimates/', 'Online estimate; "Your Pricing" step after contact info', 'NEEDS REQUEST', 'Form asks sq ft, beds, baths (whole numbers), clean level, frequency; no partial-home option on the form', 'Yes – name, email, phone, address', 'Southern Louisville office (3411 Bardstown Rd, 40218)', 'Submit (needs OK); ask about partial-home cleaning', '—', ''),
]
for i, row in enumerate(T):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=5 + i, column=j, value=v)
        c.font = N
        c.alignment = Alignment(wrap_text=True, vertical='top')
    st_ = ws.cell(row=5 + i, column=6)
    st_.font = Font(name=F, bold=True, color={'CAPTURED': '00703C', 'PUBLISHED': '1F4E78', 'NEEDS REQUEST': 'C00000'}.get(row[5], '000000'))
r0 = 6 + len(T)
ws.cell(row=r0, column=1, value='Marketing follow-up log (fill in after any request is submitted)').font = Font(name=F, bold=True, size=12)
header(ws, r0 + 1, ['Company', 'Request date', 'Marketing consent ticked?', 'First response (date/time)', 'Hours to first response', 'Quote-related msgs', 'Promotional msgs', 'Msgs within 7 days', 'Msgs within 30 days', 'Quoted price', 'Notes', ''])
ex_row = ['Example Co', '2026-10-14 10:00', 'No', '2026-10-14 13:30', None, 1, 0, 1, 1, 249, 'Example row – overwrite']
for j, v in enumerate(ex_row, 1):
    c = ws.cell(row=r0 + 2, column=j, value=v)
    c.font = Font(name=F, italic=True, color='808080')
ws.cell(row=r0 + 2, column=5, value=f'=IFERROR((DATEVALUE(LEFT(D{r0+2},10))+TIMEVALUE(RIGHT(D{r0+2},5))-DATEVALUE(LEFT(B{r0+2},10))-TIMEVALUE(RIGHT(B{r0+2},5)))*24,"")').number_format = '0.0'
for k in range(10):
    for j in range(1, 12):
        ws.cell(row=r0 + 3 + k, column=j).fill = YEL
widths(ws, [4, 30, 10, 46, 34, 15, 46, 30, 26, 44, 40, 15])
ws.freeze_panes = 'C5'

# ------------------------------------------------------------------ Platforms
ws = wb.create_sheet('Platforms', 2)
ws['A1'] = 'Independent cleaners and platforms – one row per provider (no contact made; public listings only)'
ws['A1'].font = TITLE
header(ws, 3, ['Channel', 'Provider (anonymized for individuals)', 'Advertised rate', 'Pricing basis', 'Minimum / fees', 'Supplies', 'Scope / notes', 'Reputation', 'Status', 'Source'])
P = [
 ('Taskrabbit', 'Tasker A (5.0★, 75 reviews, 96 cleaning tasks)', 'Rate shown only after entering a street address', 'Hourly', '2-hour minimum', 'Brings own supplies', '5 yrs residential/business cleaning', '5.0 (75)', 'NEEDS ADDRESS STEP', 'https://www.taskrabbit.com/locations/louisville/house-cleaning'),
 ('Taskrabbit', 'Tasker B (5.0★, 54 reviews, 65 tasks)', 'Rate shown only after entering a street address', 'Hourly', 'Deep clean of stove/fridge/baseboards billed +1 hour', 'Brings own supplies', 'Dusting, vacuuming, surfaces, mopping, trash', '5.0 (54)', 'NEEDS ADDRESS STEP', 'https://www.taskrabbit.com/locations/louisville/house-cleaning'),
 ('Taskrabbit', 'Tasker C (5.0★, 46 reviews, 63 tasks)', 'Rate shown only after entering a street address', 'Hourly', '—', '—', 'Hotel / Airbnb background; business insurance', '5.0 (46)', 'NEEDS ADDRESS STEP', 'https://www.taskrabbit.com/locations/louisville/house-cleaning'),
 ('Taskrabbit', 'Platform-wide', '"House cleaning prices begin at $40 in Louisville"', 'Hourly', 'Plus Taskrabbit service/trust fee at checkout', '—', '325 cleaning pros listed in Louisville', '—', 'PUBLISHED', 'https://www.taskrabbit.com/locations/louisville/house-cleaning'),
 ('Thumbtack', 'Cleanoria House Cleaning Louisville', 'Standard $125–$495, deep $145–$555 by size (full table on Comparison sheets)', 'Flat by sq ft band + bed/bath', 'Prices start at $100', 'Brings supplies', 'Recurring priced lower; extras extra; 2 employees', '4.8 (17), hired 59×', 'PUBLISHED', 'https://www.thumbtack.com/ky/louisville/house-cleaning/cleanoria-house-cleaning-louisville/service/569098244546895875'),
 ('Thumbtack', 'These Girls Cleaning Company', '"Roughly $60 an hour"', 'Hourly', '—', '—', '', '4.8, hired 19×', 'PUBLISHED', 'https://www.thumbtack.com/ky/louisville/house-cleaning/these-girls-cleaning-company/service/304974182379905228'),
 ('Thumbtack', 'Soclean Cleaning Services', 'Hourly for the first clean (rate not published)', 'Hourly first visit, then set price', '—', '—', 'Explains hourly avoids under-scoped flat quotes', '4.7, hired 101×', 'NEEDS REQUEST', 'https://www.thumbtack.com/ky/louisville/house-cleaning/soclean-cleaning-services/service/253644441466537196'),
 ('Thumbtack', "Christy's Commercial & Residential Cleaning", 'Hourly most popular (rate not published)', 'Hourly or flat', '—', '—', '', '4.9, hired 82×', 'NEEDS REQUEST', 'https://www.thumbtack.com/ky/louisville/house-cleaning/christys-commercial-residential-cleaning/service/267409768300618809'),
 ('Thumbtack', 'Scrub Lou (Thumbtack listing)', '"Starting at $150"', 'Custom', '—', '—', 'Recurring discounts', '4.9, hired 161×', 'PUBLISHED', 'https://www.thumbtack.com/ky/louisville/house-cleaning/scrub-lou/service/520215389162184719'),
 ('Homeaglow', 'Homeaglow (pricing flow)', 'Standard rate ≈ $25/hr ($50 for 2 hrs, $75 for 3 hrs, $100 for 4 hrs, $150 for 6 hrs); members "~60% off"', 'Hours × cleaner rate', 'Requires ForeverClean membership: $19 first month, then $59/month; 6-month term', 'Cleaner-dependent', 'Intro: free 3-hr clean (or 4 hrs for $19, 6 hrs for $59) with membership. Cancel online any time, but cancelling in first 6 months = pay full price for the first clean. Free cancellation/reschedule up to 6 hrs before.', '4.8 (4,849 site); 3.9 on Thumbtack', 'PUBLISHED', 'https://www.homeaglow.com/'),
 ('Angi', 'Matching request', '—', '—', '—', '—', 'Requires submitting a project request with contact details – not done', '—', 'NEEDS REQUEST', 'https://www.angi.com/'),
 ('Care.com', 'Cleaner A (6 yrs, hired by 9 local families)', 'from $30/hr', 'Hourly', 'Agreed total/minimum only via messaging (account needed)', '—', 'Basement, carpet, linens, baths, cabinets', '5.0 (1)', 'PUBLISHED (advertised rate only)', 'https://www.care.com/house-cleaning/louisville-ky'),
 ('Care.com', 'Cleaner B (10 yrs, hired by 5)', 'from $40/hr', 'Hourly', 'Same', '—', 'Baths, rooms, kitchen, linens, dusting', '5.0 (1)', 'PUBLISHED (advertised rate only)', 'https://www.care.com/house-cleaning/louisville-ky'),
 ('Care.com', 'Cleaner C (owns cleaning business 12 yrs)', 'from $20/hr', 'Hourly', 'Same', '—', 'Recurring, one-time, move-in/out, deep', '5.0 (1)', 'PUBLISHED (advertised rate only)', 'https://www.care.com/house-cleaning/louisville-ky'),
 ('Care.com', 'Platform-wide', 'Average posted rate $19.00/hr (Oct 2026); listings $15–$40/hr', 'Hourly', 'Care.com membership needed to message', '—', '', '—', 'PUBLISHED', 'https://www.care.com/house-cleaning/louisville-ky'),
 ('Craigslist', 'Ad 1 – local business (links to Super Cleaning Service)', '"Up to 45% off"; flat rates', 'Flat rate', '—', '—', 'Recurring, one-time, deep, move-in/out, carpet', '—', 'NO PRICE', 'https://www.craigslist.org/view/d/louisville-best-house-and-carpet/x67N43K3N9rv3AidJNwvkL'),
 ('Craigslist', 'Ad 2 – individual (posted 9/29)', '"Reasonable prices" – no number', 'Estimate', '—', '—', 'Houses & apartments; weekly, biweekly, one-time', '—', 'NO PRICE', 'https://www.craigslist.org/view/d/louisville-house-cleaning-services/mRzH52LDJ3vMJ9AAhtkE56'),
 ('Craigslist', 'Ad 3 – small LLC (posted 9/24)', '$25 off first service; free quotes', 'Quote', '—', '—', 'Standard, deep, move-out, apartment turnovers, oven/fridge, cabinets, baseboards, windows', '—', 'NO PRICE', 'https://www.craigslist.org/view/d/louisville-professional-house-cleaning/fo99WMmb55XthuunLBUSuN'),
]
for i, row in enumerate(P):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=4 + i, column=j, value=v)
        c.font = N
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws.cell(row=5 + len(P), column=1, value='Individuals are listed by profile stats rather than names. Care.com rates are "from" rates set by the cleaner; agreed totals need a Care.com account and messaging.').font = Font(name=F, italic=True)
widths(ws, [13, 38, 40, 18, 32, 16, 50, 18, 22, 60])
ws.freeze_panes = 'C4'

# Rocking Robins rules on Competitor Rules sheet
ws = wb['Competitor Rules']
r = ws.max_row + 2
ws.cell(row=r, column=1, value='ROCKING ROBINS (R&R) MAIDS — price = hours × $60; hours = layout + sq ft + condition + basic/deep + stories + basement (first clean only)').font = B
r += 1
for k, sect in [('Layout hours', RRM['bed_bath_hours']), ('Sq ft hours', RRM['sqft_hours']), ('Condition', RRM['condition_hours']), ('Cleaning type', RRM['cleaning_type_hours']), ('Stories', RRM['stories_hours']), ('Basement', RRM['basement_hours'])]:
    ws.cell(row=r, column=1, value=k).font = B
    r += 1
    for kk, vv in sect.items():
        ws.cell(row=r, column=1, value=kk).font = N
        ws.cell(row=r, column=2, value=vv).font = N
        c = ws.cell(row=r, column=3, value=f'=B{r}*60')
        c.number_format = USD0
        c.font = N
        r += 1
ws.cell(row=r, column=1, value='Add-ons: interior windows $5 each, blinds $20 each, fridge/freezer $60 each. $50 deposit at booking.').font = N

# ------------------------------------------------------------------ Full Matrix
ws = wb.create_sheet('Full Matrix')
keys = ['company', 'service', 'site_service', 'frequency', 'bedrooms', 'bathrooms', 'full_baths', 'half_baths', 'sqft_range', 'first_visit_pretax', 'recurring_visit_pretax', 'first_visit_with_tax']
labels = ['Company', 'Service', 'Site service name', 'Frequency', 'Bedrooms', 'Bathrooms', 'Full baths', 'Half baths', 'Sq ft band', 'First visit (pre-tax)', 'Ongoing visit (pre-tax)', 'First visit incl. 6% tax']
header(ws, 1, labels)
for i, row in enumerate(M['full'], 2):
    for j, k in enumerate(keys, 1):
        v = row.get(k)
        c = ws.cell(row=i, column=j, value=v)
        if j >= 10:
            c.number_format = USD
n = len(M['full']) + 1
tab = Table(displayName='FullMatrix', ref=f'A1:L{n}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
ws.add_table(tab)
widths(ws, [24, 13, 34, 14, 10, 10, 10, 10, 18, 16, 16, 16])
ws.freeze_panes = 'A2'

for s in wb.worksheets:
    if s.title != 'Full Matrix':
        fontall(s)
out = D / 'Louisville_Cleaning_Price_Matrix.xlsx'
wb.save(out)
print(out)
