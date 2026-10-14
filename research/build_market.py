"""Build the multi-city cleaning price matrix from every harvest in output/harvest/.

    python research/build_market.py            # latest harvest of every company
    python research/build_market.py --date 2026-10-20

Writes research/market/<date>/:
  market_matrix.csv     every harvested combination, normalized (same columns as 2026-10-13/full_matrix.csv + city)
  profiles.csv          every company priced on the same 19 common home profiles
  city_index.csv        each city's median price vs Louisville, per service / frequency
  structure.csv         how each company builds its price (discounts, deep/move premium, $ per bed/bath/sq ft)
  Market_Price_Matrix.xlsx

The Louisville models from research/2026-10-13 are included as a baseline until those companies are re-harvested.
Needs Python 3 with openpyxl.
"""
import csv, json, re, statistics as st, sys, argparse
from collections import defaultdict
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HARVEST = ROOT / 'output' / 'harvest'
BASELINE = ROOT / 'research' / '2026-10-13' / 'full_matrix.csv'
CITIES = {c['id']: c for c in json.load(open(ROOT / 'config' / 'cities.json'))['cities']}
CITY_BY_NAME = {c['name']: c for c in CITIES.values()}

COLS = ['company', 'city', 'state', 'zip', 'platform', 'service', 'site_service', 'frequency', 'bedrooms', 'bathrooms',
        'full_baths', 'half_baths', 'sqft_range', 'sqft_min', 'sqft_max', 'other_options', 'first_visit_pretax',
        'recurring_visit_pretax', 'first_visit_with_tax', 'tax_rate_pct', 'verified', 'captured', 'method']

PROFILES = [(1, 1, 700), (1, 1, 900), (2, 1, 1000), (2, 1.5, 1100), (2, 2, 1200), (3, 1, 1250), (3, 1.5, 1400), (3, 2, 1600),
            (3, 2, 1800), (3, 2.5, 2000), (3, 3, 2200), (4, 2, 2000), (4, 2.5, 2300), (4, 3, 2700), (4, 3.5, 3200),
            (5, 3, 3000), (5, 3.5, 3600), (5, 4, 4200), (5, 4.5, 4800)]
SVC_FREQ = [('standard', 'one_time'), ('standard', 'weekly'), ('standard', 'biweekly'), ('standard', 'every_4_weeks'),
            ('deep', 'one_time'), ('move_in_out', 'one_time')]
BASE_HOME = (3, 2, 1800)

r2 = lambda x: None if x is None else round(x + 1e-9, 2)


def money(s):
    if s in (None, ''):
        return None
    m = re.search(r'-?[\d,]+(?:\.\d+)?', str(s))
    return float(m.group(0).replace(',', '')) if m else None


def num(s):
    if s in (None, ''):
        return None
    if re.search(r'studio', str(s), re.I):
        return 0.0
    m = re.search(r'\d+(?:\.\d+)?', str(s))
    return float(m.group(0)) if m else None


def sqft_bounds(label):
    if not label:
        return None, None
    s = str(label).replace(',', '')
    n = [int(x) for x in re.findall(r'\d+', s)]
    if not n:
        return None, None
    if re.search(r'under|less|below|up to|<', s, re.I):
        return 0, n[0]
    if len(n) == 1 or re.search(r'\+|over|more|above', s, re.I):
        return n[0], 99999
    return n[0], n[1]


def norm_service(name, addon=''):
    s = f'{name} {addon}'.lower()
    if re.search(r'office|commercial|carpet|window|post.?construct|janitor|pressure|airbnb|vacation|organiz|hourly|per hour', s):
        return 'other'
    if re.search(r'move', s):
        return 'move_in_out'
    if re.search(r'deep|first.?time|initial|spring', s):
        return 'deep'
    return 'standard'


def norm_freq(f):
    s = (f or '').lower()
    if not s or s == '(none)' or re.search(r'one|once|single', s):
        return 'one_time'
    if re.search(r'2 ?weeks|other week|bi.?week|fortnight|twice a month|every two', s):
        return 'biweekly'
    if re.search(r'4 ?weeks|month|every four', s):
        return 'every_4_weeks'
    if re.search(r'twice a week|daily|2x', s):
        return 'other'
    if re.search(r'week', s):
        return 'weekly'
    return 'other'


def latest_files(day=None):
    """Newest price CSV per company (or the newest on/before --date)."""
    out = {}
    if not HARVEST.exists():
        return out
    for d in HARVEST.iterdir():
        if not d.is_dir() or d.name.startswith('_'):
            continue
        files = sorted(f for f in d.glob('*.csv') if not re.search(r'_(form|raw-feed|extras)\.csv$', f.name))
        if day:
            files = [f for f in files if f.name[:10] <= day]
        if files:
            out[d.name] = files[-1]
    return out


def read_harvest(path):
    rows = []
    with open(path, encoding='utf-8-sig') as fh:
        for r in csv.DictReader(fh):
            city = CITY_BY_NAME.get(r.get('city', ''), {})
            tax = money(r.get('taxRatePct'))
            if 'total' in r:  # BookingKoala: totals as displayed
                tbt, tot, rec = money(r.get('totalBeforeTax')), money(r.get('total')), money(r.get('recurringTotal'))
                k = 1 + (tax or 0) / 100
                pre = tbt if tbt is not None else (tot / k if tot is not None else None)
                with_tax = tot if tbt is not None or tax else None
                rec_pre = None if rec is None else rec / k
                site_service, addon, method = r.get('service', ''), '', 'live form capture (BookingKoala Booking Summary)'
                full_b = half_b = None
                other = r.get('otherOptions', '')
                verified = r.get('verified', '').lower() == 'true'
            else:  # Launch27: price feed
                pre, rec_pre = money(r.get('firstVisitPreTax')), money(r.get('recurringVisitPreTax'))
                with_tax = None if tax is None else pre * (1 + tax / 100)
                site_service, addon = r.get('service', ''), r.get('addOn', '')
                site_service = site_service + (f' + {addon}' if addon else '')
                method = "site's own Launch27 price feed"
                full_b, half_b = num(r.get('fullBaths')), num(r.get('halfBaths'))
                other = r.get('options', '')
                verified = True
            if pre is None:
                continue
            freq = norm_freq(r.get('frequency'))
            lo, hi = sqft_bounds(r.get('sqft'))
            rows.append(dict(company=r.get('company'), city=city.get('id', r.get('city', '')), state=r.get('state', ''), zip=r.get('zip', ''),
                             platform=r.get('platform', ''), service=norm_service(r.get('service', ''), addon), site_service=site_service,
                             frequency=freq, bedrooms=num(r.get('bedrooms')), bathrooms=num(r.get('bathrooms')), full_baths=full_b, half_baths=half_b,
                             sqft_range=r.get('sqft', ''), sqft_min=lo, sqft_max=hi, other_options=other,
                             first_visit_pretax=r2(pre), recurring_visit_pretax=r2(rec_pre) if freq != 'one_time' else None,
                             first_visit_with_tax=r2(with_tax), tax_rate_pct=tax, verified=verified, captured=r.get('captured', path.name[:10]), method=method))
    return rows


def read_baseline(skip_names):
    rows = []
    if not BASELINE.exists():
        return rows
    for r in csv.DictReader(open(BASELINE, encoding='utf-8-sig')):
        if any(n.lower().startswith(r['company'].lower()) or r['company'].lower().startswith(n.lower()) for n in skip_names):
            continue
        lo, hi = sqft_bounds(r['sqft_range'])
        rows.append(dict(company=r['company'], city='louisville', state='KY', zip='40217', platform='', service=r['service'],
                         site_service=r['site_service'], frequency=r['frequency'], bedrooms=num(r['bedrooms']), bathrooms=num(r['bathrooms']),
                         full_baths=num(r['full_baths']), half_baths=num(r['half_baths']), sqft_range=r['sqft_range'], sqft_min=lo, sqft_max=hi,
                         other_options='', first_visit_pretax=money(r['first_visit_pretax']), recurring_visit_pretax=money(r['recurring_visit_pretax']),
                         first_visit_with_tax=money(r['first_visit_with_tax']), tax_rate_pct=6, verified=True, captured='2026-10-13',
                         method=r['method'] + ' (2026-10-13 baseline)'))
    return rows


def build_index(rows):
    """company -> (service, frequency) -> list of rows usable for profile lookup."""
    idx = defaultdict(lambda: defaultdict(list))
    for r in rows:
        if r['service'] == 'other' or r['frequency'] == 'other':
            continue
        idx[(r['company'], r['city'])][(r['service'], r['frequency'])].append(r)
    return idx


def lookup(cands, bed, bath, sqft):
    """Exact bed + bath, sq ft range containing the home; lowest-priced default for any other options."""
    best = None
    for r in cands:
        if r['bedrooms'] is not None and r['bedrooms'] != bed:
            continue
        if r['bathrooms'] is not None and abs(r['bathrooms'] - bath) > 1e-9:
            continue
        if r['sqft_min'] is not None and not (r['sqft_min'] <= sqft <= r['sqft_max']):
            continue
        if r['bedrooms'] is None and r['bathrooms'] is None and r['sqft_min'] is None:
            continue
        if best is None or r['first_visit_pretax'] < best['first_visit_pretax']:
            best = r
    return best


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--date', help='use the newest harvest on or before this date (YYYY-MM-DD)')
    args = ap.parse_args()
    day = args.date or date.today().isoformat()
    out = ROOT / 'research' / 'market' / day
    out.mkdir(parents=True, exist_ok=True)

    rows = []
    files = latest_files(args.date)
    for cid, f in sorted(files.items()):
        got = read_harvest(f)
        print(f'  {cid:40s} {len(got):7d} rows  ({f.name})')
        rows += got
    harvested_names = {r['company'] for r in rows}
    base = read_baseline({n for n in harvested_names if n})
    if base:
        print(f'  + {len(base)} Louisville baseline rows (2026-10-13 models for companies not re-harvested)')
    rows += base
    if not rows:
        sys.exit('No harvest data found in output/harvest/. Run: npm run collect')

    with open(out / 'market_matrix.csv', 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, COLS)
        w.writeheader()
        w.writerows(rows)

    # ---- every company on the same profiles
    idx = build_index(rows)
    prof_rows = []
    for (company, city), by in sorted(idx.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        for (bed, bath, sqft) in PROFILES:
            for svc, freq in SVC_FREQ:
                hit = lookup(by.get((svc, freq), []), bed, bath, sqft)
                if hit:
                    prof_rows.append(dict(city=city, company=company, profile=f'{bed} bd / {bath:g} ba / {sqft:,} sq ft', bedrooms=bed, bathrooms=bath,
                                          sqft=sqft, service=svc, frequency=freq, first_visit_pretax=hit['first_visit_pretax'],
                                          recurring_visit_pretax=hit['recurring_visit_pretax'], matched_sqft_range=hit['sqft_range'], verified=hit['verified']))
    write(out / 'profiles.csv', prof_rows)

    # ---- city medians and index vs Louisville
    med = defaultdict(list)
    for p in prof_rows:
        med[(p['city'], p['service'], p['frequency'], p['profile'])].append(p['recurring_visit_pretax'] or p['first_visit_pretax'] if p['frequency'] != 'one_time' else p['first_visit_pretax'])
    city_rows = []
    home = next((c for c in CITIES.values() if c.get('home')), {'id': 'louisville'})['id']
    for cid, c in CITIES.items():
        row = dict(city=c['name'], state=c['state'], companies=len({k[0] for k in idx if k[1] == cid}))
        for svc, freq in SVC_FREQ:
            ratios, vals = [], []
            for (bed, bath, sqft) in PROFILES:
                pl = f'{bed} bd / {bath:g} ba / {sqft:,} sq ft'
                v = med.get((cid, svc, freq, pl))
                h = med.get((home, svc, freq, pl))
                if v:
                    vals.append(st.median(v))
                    if h:
                        ratios.append(st.median(v) / st.median(h))
            bl = med.get((cid, svc, freq, f'{BASE_HOME[0]} bd / {BASE_HOME[1]:g} ba / {BASE_HOME[2]:,} sq ft'))
            row[f'{svc} {freq}: median 3bd/2ba/1800'] = r2(st.median(bl)) if bl else None
            row[f'{svc} {freq}: index vs {CITIES[home]["name"]}'] = round(st.median(ratios), 3) if ratios else None
        city_rows.append(row)
    write(out / 'city_index.csv', city_rows)

    # ---- pricing structure per company (what to copy, not just the dollar level)
    struct = []
    for (company, city), by in sorted(idx.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        def p(svc, freq, bed, bath, sqft, rec=False):
            h = lookup(by.get((svc, freq), []), bed, bath, sqft)
            if not h:
                return None
            return h['recurring_visit_pretax'] if rec and h['recurring_visit_pretax'] is not None else h['first_visit_pretax']
        b = p('standard', 'one_time', *BASE_HOME)
        row = dict(city=CITIES.get(city, {}).get('name', city), company=company, rows=sum(len(v) for v in by.values()),
                   **{'standard one-time 3bd/2ba/1800': b})
        for freq in ['weekly', 'biweekly', 'every_4_weeks']:
            f1, fr = p('standard', freq, *BASE_HOME), p('standard', freq, *BASE_HOME, rec=True)
            row[f'{freq} first-visit discount'] = r2(1 - f1 / b) if b and f1 else None
            row[f'{freq} recurring discount'] = r2(1 - fr / b) if b and fr else None
        for svc in ['deep', 'move_in_out']:
            v = p(svc, 'one_time', *BASE_HOME)
            row[f'{svc} premium x standard'] = r2(v / b) if b and v else None
        pb4, pb2 = p('standard', 'one_time', 4, 2, 1800), p('standard', 'one_time', 2, 2, 1800)
        row['$ per extra bedroom'] = r2((pb4 - pb2) / 2) if pb4 and pb2 else None
        pa3, pa1 = p('standard', 'one_time', 3, 3, 1800), p('standard', 'one_time', 3, 1, 1800)
        row['$ per extra full bath'] = r2((pa3 - pa1) / 2) if pa3 and pa1 else None
        s1, s3 = p('standard', 'one_time', 3, 2, 1200), p('standard', 'one_time', 3, 2, 3200)
        row['$ per extra 500 sq ft'] = r2((s3 - s1) / 4) if s1 and s3 else None
        struct.append(row)
    write(out / 'structure.csv', struct)

    try:
        workbook(out, rows, prof_rows, city_rows, struct)
    except ImportError:
        print('openpyxl not installed: CSVs written, workbook skipped (pip install openpyxl)')
    print(f'\n{len(rows):,} rows from {len(idx)} companies in {len({k[1] for k in idx})} cities -> {out}')


def write(path, rows):
    if not rows:
        return
    with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def workbook(out, rows, prof_rows, city_rows, struct):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    F = 'Arial'
    H, HF = Font(name=F, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
    USD, PCT, X = '$#,##0.00;($#,##0.00);-', '0.0%', '0.00"x"'
    wb = Workbook()

    def sheet(ws, data, money=(), pct=(), ratio=(), width=14):
        if not data:
            ws['A1'] = 'No data yet'
            return
        keys = list(data[0].keys())
        for i, k in enumerate(keys, 1):
            c = ws.cell(row=1, column=i, value=k)
            c.font, c.fill, c.alignment = H, HF, Alignment(wrap_text=True, vertical='center', horizontal='center')
            ws.column_dimensions[get_column_letter(i)].width = max(width, min(42, len(k) * 0.9)) if i > 2 else 26
        for r, d in enumerate(data, 2):
            for i, k in enumerate(keys, 1):
                c = ws.cell(row=r, column=i, value=d[k])
                c.font = Font(name=F)
                if any(m in k for m in money) or k in money:
                    c.number_format = USD
                if any(m in k for m in pct):
                    c.number_format = PCT
                if any(m in k for m in ratio):
                    c.number_format = X
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 48

    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = 'Residential Cleaning – Multi-City Competitor Price Matrix'
    ws['A1'].font = Font(name=F, bold=True, size=14)
    cities = sorted({r['city'] for r in rows})
    lines = [
        f'Built {date.today().isoformat()} from {len({(r["company"], r["city"]) for r in rows})} companies in {len(cities)} cities, {len(rows):,} priced combinations. Prices are pre-tax.',
        'City Index: each city\'s median price on the same 19 home profiles, divided by Louisville\'s. Use it to translate peer-city prices into Louisville terms.',
        'Pricing Structure: how each company builds its price (frequency discounts, deep / move-out premium, $ per bedroom, bath and 500 sq ft) on a 3 bd / 2 ba / 1,800 sq ft home.',
        'Profiles: every company priced on the same 19 home profiles x 6 service/frequency cases.',
        'Full Matrix: every combination each company\'s online form offers (same columns as the 13 Oct full_matrix.csv, plus city).',
        'Only companies with instant online prices are included (BookingKoala and Launch27 forms). No inquiries or bookings were submitted.',
    ]
    for i, l in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=l).font = Font(name=F)
    ws.column_dimensions['A'].width = 140
    sheet(wb.create_sheet('City Index'), city_rows, money=('median',), ratio=('index',))
    sheet(wb.create_sheet('Pricing Structure'), struct, money=('3bd/2ba', '$ per'), pct=('discount',), ratio=('premium',))
    sheet(wb.create_sheet('Profiles'), prof_rows, money=('pretax',))
    ws = wb.create_sheet('Full Matrix')
    ws.append(COLS)
    for c in ws[1]:
        c.font, c.fill = H, HF
    for r in rows[:1_000_000]:
        ws.append([r[k] for k in COLS])
    ws.freeze_panes = 'A2'
    runs = sorted((HARVEST / '_runs').glob('*.csv')) if (HARVEST / '_runs').exists() else []
    if runs:
        log = list(csv.DictReader(open(runs[-1], encoding='utf-8-sig')))
        sheet(wb.create_sheet('Harvest Log'), log)
    disc = sorted((ROOT / 'output' / 'discovery').glob('*.csv')) if (ROOT / 'output' / 'discovery').exists() else []
    if disc:
        seen, dl = set(), []
        for f in disc:
            for d in csv.DictReader(open(f, encoding='utf-8-sig')):
                if d['domain'] not in seen:
                    seen.add(d['domain'])
                    dl.append(d)
        sheet(wb.create_sheet('Companies Checked'), dl)
    wb.save(out / 'Market_Price_Matrix.xlsx')


if __name__ == '__main__':
    main()
